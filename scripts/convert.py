"""知识库入库辅助：把 PDF / 表格转成 Markdown 写到 stdout。

用法::

    python scripts/convert.py pdf  <file>
    python scripts/convert.py xlsx <file>
    python scripts/convert.py csv  <file>

只依赖本机已装的 pdfplumber / openpyxl；缺失时给出安装提示并以非零码退出，
由调用方（knowledgeAdd.ts）决定是跳过还是中断。
"""

import csv
import sys


def rows_to_md(rows):
    """把二维表转成 Markdown 表格；空行会被丢掉。"""
    rows = [list(r) for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        return ""

    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    header, body = rows[0], rows[1:]

    lines = ["| " + " | ".join(header) + " |", "|" + "---|" * width]
    for row in body:
        cells = [str(c).replace("\n", " ").replace("|", r"\|") for c in row]
        lines.append("| " + " | ".join(cells) + " |")
    return "\n".join(lines)


def pdf_to_md(path):
    try:
        import pdfplumber
    except ImportError:
        sys.exit("缺少 pdfplumber，请先执行: python -m pip install pdfplumber")

    parts = []
    with pdfplumber.open(path) as pdf:
        for index, page in enumerate(pdf.pages, 1):
            text = (page.extract_text() or "").strip()
            if text:
                # 逐页加标题：既是切分边界，也让回答能引用到具体页
                parts.append("## 第 %d 页\n\n%s" % (index, text))
    return "\n\n".join(parts)


def xlsx_to_md(path):
    try:
        from openpyxl import load_workbook
    except ImportError:
        sys.exit("缺少 openpyxl，请先执行: python -m pip install openpyxl")

    parts = []
    workbook = load_workbook(path, data_only=True)
    for sheet in workbook.worksheets:
        rows = [
            ["" if cell is None else cell for cell in row]
            for row in sheet.iter_rows(values_only=True)
        ]
        table = rows_to_md(rows)
        if table:
            parts.append("## %s\n\n%s" % (sheet.title, table))
    return "\n\n".join(parts)


def csv_to_md(path):
    # utf-8-sig 兼容 Excel 导出的 BOM 头
    with open(path, newline="", encoding="utf-8-sig") as handle:
        return rows_to_md(list(csv.reader(handle)))


CONVERTERS = {"pdf": pdf_to_md, "xlsx": xlsx_to_md, "csv": csv_to_md}


def main():
    # Windows 下 stdout 默认走本地编码，中文会乱码或抛错
    sys.stdout.reconfigure(encoding="utf-8")

    if len(sys.argv) != 3 or sys.argv[1] not in CONVERTERS:
        sys.exit(__doc__)

    markdown = CONVERTERS[sys.argv[1]](sys.argv[2])
    if not markdown.strip():
        sys.exit("未能从 %s 提取到任何文本（可能是扫描件或纯图片表格）" % sys.argv[2])

    sys.stdout.write(markdown)


if __name__ == "__main__":
    main()
